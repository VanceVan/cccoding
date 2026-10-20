# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=(0x1F,0x3B,0x73); AVAIL=6.7
def newdoc(margin=1.0, size=11.0):
    doc=Document(); st=doc.styles['Normal']
    st.font.name='Times New Roman'; st.font.size=Pt(size)
    st.element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st.element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
    st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.08
    for s in doc.sections:
        s.top_margin=s.bottom_margin=Inches(margin); s.left_margin=s.right_margin=Inches(margin)
    return doc
def runs(p,t,size=None,font='Times New Roman'):
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',t):
        if not tok: continue
        if tok.startswith('**'): r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('*'): r=p.add_run(tok[1:-1]); r.italic=True
        else: r=p.add_run(tok)
        r.font.name=font
        if size: r.font.size=Pt(size)
        r._element.rPr.rFonts.set(qn('w:ascii'),font); r._element.rPr.rFonts.set(qn('w:hAnsi'),font)
    return p
def mk(doc):
    def P(t,size=None,align=None,after=7,justify=True,font='Times New Roman'):
        p=doc.add_paragraph(); runs(p,t,size,font); p.paragraph_format.space_after=Pt(after)
        if justify and not align: p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if align: p.alignment=align
        return p
    def H(t,size=12,before=11,color=NAVY,font='Times New Roman'):
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_with_next=True
        r=p.add_run(t); r.bold=True; r.font.size=Pt(size); r.font.name=font; r.font.color.rgb=RGBColor(*color)
        r._element.rPr.rFonts.set(qn('w:ascii'),font); r._element.rPr.rFonts.set(qn('w:hAnsi'),font)
        return p
    def shade(c,hx):
        tcPr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hx); tcPr.append(sh)
    def TABLE(hdr,rows,widths,fs=8.6,font='Times New Roman'):
        tot=sum(widths); widths=[w*AVAIL/tot for w in widths]
        t=doc.add_table(rows=1,cols=len(hdr)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        tblPr=t._tbl.tblPr
        mar=OxmlElement('w:tblCellMar')
        for side,v in (('left','50'),('right','50'),('top','20'),('bottom','20')):
            e=OxmlElement('w:'+side); e.set(qn('w:w'),v); e.set(qn('w:type'),'dxa'); mar.append(e)
        tblPr.append(mar)
        lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
        g=t._tbl.find(qn('w:tblGrid'))
        if g is not None: t._tbl.remove(g)
        g=OxmlElement('w:tblGrid')
        for w in widths:
            gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(int(w*1440))); g.append(gc)
        tblPr.addnext(g)
        for i,hh in enumerate(hdr):
            c=t.rows[0].cells[i]; c.text=''; pp=c.paragraphs[0]
            r=pp.add_run(hh); r.bold=True; r.font.size=Pt(fs); r.font.name=font
            pp.paragraph_format.space_after=Pt(1); shade(c,'E8EAED')
        for row in rows:
            cs=t.add_row().cells
            for i,v in enumerate(row):
                c=cs[i]; c.text=''; pp=c.paragraphs[0]; runs(pp,v,fs,font); pp.paragraph_format.space_after=Pt(1)
        for r_ in t.rows:
            for i,w in enumerate(widths): r_.cells[i].width=Inches(w)
        trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
        for r_ in t.rows:
            cant=OxmlElement('w:cantSplit'); cant.set(qn('w:val'),'true'); r_._tr.get_or_add_trPr().append(cant)
        for c_ in t.rows[-1].cells:
            for pp_ in c_.paragraphs: pp_.paragraph_format.keep_with_next=True
        return t
    def FIG(path,width,cap):
        doc.add_picture(path,width=Inches(width))
        doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_before=Pt(6)
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        return P(cap,size=9,align=WD_ALIGN_PARAGRAPH.CENTER,after=10,justify=False)
    return P,H,TABLE,FIG
